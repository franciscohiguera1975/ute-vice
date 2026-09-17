"""Exportador a Excel (.xlsx)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.enums import FormatoReporte
from app.domain.ports.reportes import ArchivoReporte, TablaReporte
from app.infrastructure.reportes.formato import (
    a_local,
    nombre_archivo,
    valor_nativo,
)

# Paleta institucional sobria: azul para cabeceras, gris para bandas.
_AZUL = "1F3864"
_AZUL_CLARO = "D9E2F3"
_GRIS = "F2F2F2"

_ALINEACION = {"izquierda": "left", "centro": "center", "derecha": "right"}


class ExportadorExcel:
    """Implementacion del puerto `ExportadorReporte` para hojas de calculo.

    Produce un archivo listo para trabajar: cabecera institucional, constancia
    de los filtros aplicados, panel inmovilizado, autofiltro y columnas
    dimensionadas. Los numeros y fechas van con su tipo nativo para que se puedan
    ordenar y sumar.
    """

    @property
    def formato(self) -> FormatoReporte:
        return FormatoReporte.XLSX

    @property
    def tipo_mime(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    @property
    def extension(self) -> str:
        return "xlsx"

    def exportar(self, tabla: TablaReporte) -> ArchivoReporte:
        libro = Workbook()
        hoja = libro.active
        assert hoja is not None
        hoja.title = "Reporte"

        # `solo_datos` deja la cabecera en la fila 1: es lo que permite contrastar
        # el archivo contra el original del que salieron los datos.
        inicio = 1 if tabla.solo_datos else self._escribir_cabecera(hoja, tabla)
        fila = self._escribir_tabla(hoja, tabla, inicio)
        if not tabla.solo_datos:
            self._escribir_totales(hoja, tabla, fila)
        self._ajustar(hoja, tabla)

        buffer = BytesIO()
        libro.save(buffer)
        return ArchivoReporte(
            nombre_archivo=nombre_archivo(tabla.titulo, self.extension, tabla.generado_en),
            contenido=buffer.getvalue(),
            tipo_mime=self.tipo_mime,
            formato=self.formato,
        )

    # ------------------------------------------------------------ secciones
    def _escribir_cabecera(self, hoja: Worksheet, tabla: TablaReporte) -> int:
        ancho = max(len(tabla.columnas), 1)

        hoja.cell(row=1, column=1, value="UNIVERSIDAD TECNOLOGICA EQUINOCCIAL").font = Font(
            bold=True, size=13, color=_AZUL
        )
        hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)

        hoja.cell(row=2, column=1, value=tabla.titulo).font = Font(bold=True, size=11)
        hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho)

        fila = 3
        if tabla.subtitulo:
            hoja.cell(row=fila, column=1, value=tabla.subtitulo).font = Font(
                italic=True, size=10, color="595959"
            )
            hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
            fila += 1

        marca = a_local(tabla.generado_en).strftime("%d/%m/%Y %H:%M")
        pie = f"Generado: {marca}"
        if tabla.generado_por:
            pie += f"  ·  Por: {tabla.generado_por}"
        pie += f"  ·  Registros: {tabla.total_filas:,}".replace(",", ".")
        hoja.cell(row=fila, column=1, value=pie).font = Font(size=9, color="808080")
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
        fila += 1

        if tabla.filtros_aplicados:
            texto = "  ·  ".join(f"{k}: {v}" for k, v in tabla.filtros_aplicados.items())
            celda = hoja.cell(row=fila, column=1, value=f"Filtros aplicados — {texto}")
            celda.font = Font(size=9, italic=True, color="806000")
            celda.fill = PatternFill("solid", fgColor="FFF2CC")
            hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
            fila += 1

        return fila + 1

    def _escribir_tabla(self, hoja: Worksheet, tabla: TablaReporte, inicio: int) -> int:
        borde = Side(style="thin", color="BFBFBF")
        marco = Border(left=borde, right=borde, top=borde, bottom=borde)

        for indice, columna in enumerate(tabla.columnas, start=1):
            celda = hoja.cell(row=inicio, column=indice, value=columna.titulo)
            # Una columna puede pedir su propio color de cabecera para reproducir
            # una plantilla institucional; sobre fondo claro el texto blanco no
            # se leeria, asi que la letra se ajusta al fondo.
            fondo = columna.color_cabecera or _AZUL
            celda.font = Font(bold=True, size=10, color="FFFFFF" if _es_oscuro(fondo) else "000000")
            celda.fill = PatternFill("solid", fgColor=fondo)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            celda.border = marco

        fila = inicio + 1
        for indice_fila, datos in enumerate(tabla.filas):
            banda = (
                PatternFill("solid", fgColor=tabla.color_banda or _GRIS)
                if indice_fila % 2
                else None
            )
            for indice_col, columna in enumerate(tabla.columnas, start=1):
                celda = hoja.cell(
                    row=fila,
                    column=indice_col,
                    value=valor_nativo(datos.get(columna.clave), columna.tipo),
                )
                celda.alignment = Alignment(
                    horizontal=_ALINEACION.get(columna.alineacion, "left"),
                    vertical="top",
                    wrap_text=columna.ancho > 30,
                )
                celda.border = marco
                celda.font = Font(size=10)
                if banda:
                    celda.fill = banda
                if columna.tipo in {"fecha", "fecha_hora"}:
                    celda.number_format = (
                        "DD/MM/YYYY" if columna.tipo == "fecha" else "DD/MM/YYYY HH:MM"
                    )
                elif columna.tipo == "numero":
                    celda.number_format = "#,##0"
                elif columna.tipo == "porcentaje":
                    celda.number_format = "0.00\\%"
            fila += 1

        # Inmoviliza cabecera y activa el autofiltro sobre el rango de datos.
        hoja.freeze_panes = hoja.cell(row=inicio + 1, column=1)
        if tabla.filas:
            ultima = get_column_letter(len(tabla.columnas))
            hoja.auto_filter.ref = f"A{inicio}:{ultima}{fila - 1}"

        return fila

    def _escribir_totales(self, hoja: Worksheet, tabla: TablaReporte, fila: int) -> None:
        if not tabla.totales:
            return
        fila += 1
        for etiqueta, valor in tabla.totales.items():
            hoja.cell(row=fila, column=1, value=f"{etiqueta}:").font = Font(bold=True, size=10)
            celda = hoja.cell(row=fila, column=2, value=valor)
            celda.font = Font(bold=True, size=10)
            celda.fill = PatternFill("solid", fgColor=_AZUL_CLARO)
            fila += 1

    def _ajustar(self, hoja: Worksheet, tabla: TablaReporte) -> None:
        for indice, columna in enumerate(tabla.columnas, start=1):
            hoja.column_dimensions[get_column_letter(indice)].width = min(
                max(columna.ancho, len(columna.titulo) + 2), 60
            )


def _es_oscuro(color: str) -> bool:
    """Luminancia percibida del color, para decidir el color de la letra."""
    try:
        r, g, b = (int(color[i : i + 2], 16) for i in (0, 2, 4))
    except (ValueError, IndexError):
        return True
    return (0.299 * r + 0.587 * g + 0.114 * b) < 140
