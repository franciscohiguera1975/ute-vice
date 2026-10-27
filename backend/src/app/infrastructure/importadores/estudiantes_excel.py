"""Lectura del reporte de matricula estudiantil de Banner.

El archivo trae una sola hoja, `RptEstBannerMatriculas`, con 27 columnas fijas.
A diferencia del distributivo docente, Banner no exporta un archivo por
periodo: el reporte trae el padron completo vigente, con `PERIODO` como una
columna mas de cada fila.

`PERIODO` es un codigo de Banner de seis digitos (`262651`). **No es el mismo
esquema que `PeriodoAcademico`/`cat_paos`** del distributivo —ese usa el
digito de nivel `15`/`65`/`75`, aqui no calzan—, asi que se guarda tal cual,
como texto plano, sin decodificarlo ni relacionarlo con ese catalogo.

`CODIGO_BECA`, `FINANCIAMIENTO`, `TIPO_DE_BECA` y `PORCENTAJE` pueden traer
varios valores separados por coma cuando un estudiante acumula mas de una
beca o descuento. Se separan aqui y se emparejan por posicion con
`itertools.zip_longest`: si una columna trae menos valores que otra, el resto
se completa con `None` en lugar de fallar.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from itertools import zip_longest
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.logging import get_logger
from app.domain.errors import ErrorValidacion
from app.domain.ports.importacion_estudiantes import BecaCruda, FilaCrudaEstudiante

log = get_logger(__name__)

#: Nombre de la unica hoja del reporte. Si no aparece -o viene renombrada- se
#: usa la primera, igual que en el lector del distributivo oficial.
_HOJA_MATRICULAS = "RptEstBannerMatriculas"

#: Columnas del reporte que se leen, con su nombre exacto en el archivo.
_COLUMNAS: dict[str, str] = {
    "sede": "SEDE",
    "facultad": "FACULTAD",
    "carrera": "CARRERA",
    "tipo_carrera": "TIPO_CARRERA",
    "modalidad": "MODALIDAD",
    "nivel_formacion_1": "NIVEL_FORMACIÓN_1",
    "nivel_formacion_2": "NIVEL_FORMACIÓN_2",
    "periodo": "PERIODO",
    "cedula": "CEDULA",
    "id_banner": "ID_BANNER",
    "nombre": "NOMBRE",
    "genero": "GÉNERO",
    "mail_personal": "MAIL_PERSONAL",
    "mail_ute": "MAIL_UTE",
    "telefono_celular": "TELEFONO_CELULAR",
    "telefono_fijo": "TELEFONO_FIJO",
    "fecha_nacimiento": "FECHA_NACIMIENTO",
    "codigo_beca": "CODIGO_BECA",
    "financiamiento": "FINANCIAMIENTO",
    "tipo_beca": "TIPO_DE_BECA",
    "porcentaje": "PORCENTAJE",
    "tipo_estudiante": "ESTUDIANTE",
    "sistema": "SISTEMA",
    "matriculado": "MATRICULADO",
    "solicitud_retiro": "SOLICITUD_DE_RETIRO",
    "bloqueo_plataforma": "BLOQUEO_PLATAFORMA",
    "estado": "ESTADO",
}

#: Sin estas no se puede construir una fila.
_OBLIGATORIAS = ("CEDULA", "PERIODO")

_AUSENTES = {"", "NA", "N/A"}


def _texto(valor: Any) -> str | None:
    if valor is None:
        return None
    limpio = " ".join(str(valor).split())
    return limpio if limpio and limpio.upper() not in _AUSENTES else None


def _lista(valor: Any) -> list[str]:
    """Parte una celda que puede traer varios valores separados por comas."""
    limpio = _texto(valor)
    if not limpio:
        return []
    return [p.strip() for p in limpio.split(",") if p.strip()]


def _fecha(valor: Any) -> date | None:
    """`DD/MM/AAAA` -> `date`. Tolera vacio o formato invalido: `None`.

    El origen la declara como texto y no como fecha nativa de Excel, pero se
    admite tambien un `datetime` por si algun libro la trae ya convertida.
    """
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor

    texto = _texto(valor)
    if not texto:
        return None
    try:
        return datetime.strptime(texto, "%d/%m/%Y").date()
    except ValueError:
        return None


def _booleano(valor: Any) -> bool:
    return (_texto(valor) or "").upper() in {"SI", "SÍ", "TRUE", "1"}


def _becas(codigo: Any, financiamiento: Any, tipo: Any, porcentaje: Any) -> tuple[BecaCruda, ...]:
    """Separa las cuatro columnas de beca por coma y las empareja por posicion.

    Si las cuatro vienen vacias, no hay ninguna beca que reportar. Si alguna
    lista trae menos elementos que otra, se completa con `None` en lugar de
    fallar: es lo que exige el negocio, porque no siempre las cuatro columnas
    estan igual de completas.
    """
    codigos = _lista(codigo)
    financiamientos = _lista(financiamiento)
    tipos = _lista(tipo)
    porcentajes = _lista(porcentaje)

    if not (codigos or financiamientos or tipos or porcentajes):
        return ()

    return tuple(
        BecaCruda(codigo_beca=c, financiamiento=f, tipo_beca=t, porcentaje=p)
        for c, f, t, p in zip_longest(codigos, financiamientos, tipos, porcentajes, fillvalue=None)
    )


class LectorEstudiantesExcel:
    """Convierte el reporte de matricula estudiantil de Banner en filas planas."""

    def leer(
        self, origen: str | Path | bytes, *, hoja: str | None = None
    ) -> list[FilaCrudaEstudiante]:
        """`origen` es una ruta o el contenido del archivo.

        Acepta bytes para que la carga desde la interfaz no tenga que
        escribir un temporal en disco, igual que `LectorPaoExcel`.
        """
        libro = self._abrir(origen)
        try:
            pagina = libro[self._elegir_hoja(libro.sheetnames, hoja)]
            iterador = pagina.iter_rows(values_only=True)

            try:
                cabecera = [str(c).strip() if c is not None else "" for c in next(iterador)]
            except StopIteration:
                raise ErrorValidacion("El archivo esta vacio", campo="archivo") from None

            indices = {nombre: i for i, nombre in enumerate(cabecera)}
            faltantes = [c for c in _OBLIGATORIAS if c not in indices]
            if faltantes:
                raise ErrorValidacion(
                    f"Al archivo le faltan columnas obligatorias: {', '.join(faltantes)}. "
                    f"Se esperaba la estructura del reporte de matricula estudiantil.",
                    campo="archivo",
                )

            filas: list[FilaCrudaEstudiante] = []
            for numero, fila in enumerate(iterador, start=2):
                cruda = self._construir(numero, fila, indices)
                if cruda is not None:
                    filas.append(cruda)

            log.info("Matricula estudiantil leida", extra={"filas": len(filas)})
            return filas
        finally:
            libro.close()

    @staticmethod
    def _abrir(origen: str | Path | bytes) -> Any:
        if isinstance(origen, bytes):
            return load_workbook(BytesIO(origen), read_only=True, data_only=True)

        camino = Path(origen)
        if not camino.exists():
            raise ErrorValidacion(f"El archivo no existe: {camino}", campo="archivo")
        return load_workbook(camino, read_only=True, data_only=True)

    @staticmethod
    def _elegir_hoja(nombres: list[str], hoja: str | None) -> str:
        if hoja:
            return hoja
        for nombre in nombres:
            if nombre.strip().lower() == _HOJA_MATRICULAS.lower():
                return nombre
        return nombres[0]

    def _construir(
        self, numero: int, fila: tuple[Any, ...], indices: dict[str, int]
    ) -> FilaCrudaEstudiante | None:
        def leer(clave: str) -> Any:
            i = indices.get(_COLUMNAS[clave])
            return fila[i] if i is not None and i < len(fila) else None

        cedula = _texto(leer("cedula"))
        periodo = _texto(leer("periodo"))
        if not cedula and not periodo:
            # Fila en blanco al final del libro: se ignora en silencio.
            return None

        return FilaCrudaEstudiante(
            numero_fila=numero,
            periodo=periodo or "",
            cedula=cedula or "",
            sede=_texto(leer("sede")),
            facultad=_texto(leer("facultad")),
            carrera=_texto(leer("carrera")),
            tipo_carrera=_texto(leer("tipo_carrera")),
            modalidad=_texto(leer("modalidad")),
            nivel_formacion_1=_texto(leer("nivel_formacion_1")),
            nivel_formacion_2=_texto(leer("nivel_formacion_2")),
            id_banner=_texto(leer("id_banner")),
            nombre=_texto(leer("nombre")),
            genero=_texto(leer("genero")),
            mail_personal=_texto(leer("mail_personal")),
            mail_ute=_texto(leer("mail_ute")),
            telefono_celular=_texto(leer("telefono_celular")),
            telefono_fijo=_texto(leer("telefono_fijo")),
            fecha_nacimiento=_fecha(leer("fecha_nacimiento")),
            becas=_becas(
                leer("codigo_beca"), leer("financiamiento"), leer("tipo_beca"), leer("porcentaje")
            ),
            tipo_estudiante=_texto(leer("tipo_estudiante")),
            sistema=_texto(leer("sistema")),
            matriculado=_texto(leer("matriculado")),
            solicitud_retiro=_texto(leer("solicitud_retiro")),
            bloqueo_plataforma=_booleano(leer("bloqueo_plataforma")),
            estado=_texto(leer("estado")),
        )
