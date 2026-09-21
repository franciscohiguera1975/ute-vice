"""Lectura del distributivo tal como lo exporta el sistema academico.

Es un formato distinto del consolidado: 74 o 75 columnas, cabeceras propias, y
la sede, el nivel y la modalidad en columnas separadas en lugar de concatenadas
dentro del nombre de la carrera.

Aporta seis datos que el consolidado no tenia —estado de validacion, fase,
semanas, relacion laboral y las dos marcas de tutoria— y **no trae el PAO**: el
periodo lo indica quien importa, porque el archivo cubre uno solo.

## Dos presentaciones de la misma hoja «Distributivo»

El sistema academico exporta la hoja sola —cabecera en la primera fila— o como
parte del reporte completo del PAO, con las demas pestanas del informe:
`Antecedentes`, `Distribucion horaria`, `Excepcionalidades`... En ese reporte
la hoja `Distributivo` no es la primera, asi que sin `hoja` de por medio no se
adivina por posicion: se busca por nombre y, si no aparece, se usa la primera
—lo que hacen los archivos que solo traen esa pestana—. Ahi la cabecera baja
un par de filas (antes va un titulo) y se parte en dos niveles:
`Titularidad`, `Categoria`, `Dedicacion` y `Relacion Laboral` aparecen dos
veces, agrupadas bajo `Antigua` y `Nueva`. Las dos dicen lo mismo salvo por
espacios sueltos, pero donde difieren de verdad `Nueva` trae el dato que
`Antigua` deja en `N/A`, asi que es la que se usa; `Antigua` se descarta
entera. La cabecera no se asume en una fila fija: se busca la primera fila
que trae `Identificacion`, `Facultad` y `Carrera/Programa`, agrupando con la
fila siguiente las columnas que la traen partida.

## El nombre de la carrera

El consolidado escribe `UIO:MEDICINA - GRADO - PRESENCIAL`. Aqui llegan las
cuatro piezas por separado y se recomponen en esa misma forma, para que una
carrera cargada por las dos vias sea la misma y no dos.

## Filas que agregan varias carreras

El origen une con comas cuando un docente dicta en varias:

    Carrera/Programa = ALIMENTOS, ELECTROMECANICA, INGENIERIA INDUSTRIAL
    Sede             = SEDE QUITO, SEDE SANTO DOMINGO

La clave natural de una fila es *una* carrera y *una* sede, y las horas vienen
como un total unico que no se puede repartir entre ellas. Esas filas **se
rechazan** en lugar de inventar una atribucion o una carrera que no existe.
"""

from __future__ import annotations

from collections.abc import Iterator, Sequence
from io import BytesIO
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook

from app.core.logging import get_logger
from app.domain.errors import ErrorValidacion
from app.domain.ports.importacion import DatosSistemaAcademico, FilaCrudaDistributivo
from app.domain.value_objects_distributivo import DistribucionHoras

log = get_logger(__name__)

#: Nombre de la facultad en el sistema academico → codigo del consolidado. Sin
#: esto cada facultad entraria dos veces, con su nombre largo y con su sigla.
FACULTADES: dict[str, str] = {
    "CIENCIAS DE LA SALUD EUGENIO ESPEJO": "FCSEE",
    "POSGRADOS EN LÍNEA": "PEL",
    "CIENCIAS DE LA INGENIERÍA E INDUSTRIAS": "FCII",
    "DERECHO, CIENCIAS ADMINISTRATIVAS Y SOCIALES": "FDCAS",
    "MEDICINA VETERINARIA Y AGRONOMÍA": "FMVA",
    "ARQUITECTURA Y URBANISMO": "FAU",
    "CIENCIAS GASTRONÓMICAS Y TURISMO": "FCGT",
    "UNIDAD ACADÉMICA ESPECIALIZADA EN LA FORMACIÓN TÉCNICA Y TECNOLÓGICA": "UAEFTT",
    "CIENCIAS, INGENIERÍA Y CONSTRUCCIÓN": "FCIC",
    "CENTRO DE EDUCACIÓN EN LÍNEA": "CEL",
    "ODONTOLOGÍA": "FO",
}

#: Sede del sistema academico → prefijo con que el consolidado nombra la
#: carrera. `UIO:MEDICINA - GRADO - PRESENCIAL`.
PREFIJO_SEDE: dict[str, str] = {
    "SEDE QUITO": "UIO",
    "SEDE SANTO DOMINGO": "STO",
    "CAMPUS CUENCA": "CUE",
    "RICARDO HIDALGO OTTOLENGHI": "MON",
}

#: Categoria del sistema academico → la del consolidado. El ERP antepone
#: «TITULAR» a las tres categorias del escalafon; el consolidado no. Sin esta
#: traduccion cada una entraria dos veces y un conteo por categoria saldria
#: partido en dos grupos donde hay uno.
CATEGORIAS: dict[str, str] = {
    "TITULAR AUXILIAR": "AUXILIAR",
    "TITULAR AUXILIAR 1": "AUXILIAR",
    "TITULAR AGREGADO": "AGREGADO",
    "TITULAR PRINCIPAL": "PRINCIPAL",
}

#: El texto del estado tal como lo escribe el origen → valor del enum.
ESTADOS: dict[str, str] = {
    "OK": "OK",
    "OK, EXCEPCIÓN": "OK_EXCEPCION",
    "OK, EXCEPCION": "OK_EXCEPCION",
    "VALIDACIÓN PENDIENTE": "PENDIENTE",
    "VALIDACION PENDIENTE": "PENDIENTE",
    "ERROR": "ERROR",
}

_CLAVES_HORAS = (
    DistribucionHoras.CLAVES_DOCENCIA
    + DistribucionHoras.CLAVES_GESTION
    + DistribucionHoras.CLAVES_INVESTIGACION
    + DistribucionHoras.CLAVES_VINCULACION
)

#: Lo que el origen escribe cuando no hay dato.
_AUSENTES = {"", "N/A", "NA", "0", "-"}


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return " ".join(str(valor).split())


def _o_nada(valor: Any) -> str | None:
    limpio = _texto(valor).upper()
    return None if limpio in _AUSENTES else _texto(valor)


#: Columnas sin las que no se puede ubicar la cabecera ni construir una fila.
_OBLIGATORIAS = ("Identificación", "Facultad", "Carrera/Programa")

#: Cuantas filas se recorren buscando la cabecera antes de rendirse. El reporte
#: completo del PAO trae dos filas de titulo antes; un margen generoso no pesa
#: en un archivo de un solo periodo.
_TOPE_FILAS_CABECERA = 20

#: Los grupos con que el reporte completo parte en dos la fila de cabecera.
#: `NUEVA` es el que se usa: ver la nota del modulo.
_GRUPOS_CABECERA = {"ANTIGUA", "NUEVA"}
_GRUPO_VIGENTE = "NUEVA"


def _combinar_cabecera(
    primaria: tuple[Any, ...], secundaria: tuple[Any, ...]
) -> tuple[dict[str, int], bool]:
    """Arma el mapa nombre de columna → posicion de una fila de cabecera.

    En la hoja suelta cada columna trae su nombre directo en `primaria`. En el
    reporte completo, un tramo de columnas queda bajo un rotulo de grupo
    (`Antigua` o `Nueva`) y el nombre real esta en `secundaria`, una fila mas
    abajo — la primera columna del tramo lleva el rotulo del grupo *y* su
    propio nombre a la vez, una encima de la otra. Se descarta el tramo
    `Antigua` entero: ver la nota del modulo.

    Devuelve tambien si se uso `secundaria`: si es asi, esa fila era parte de
    la cabecera y no un dato.
    """
    columnas: dict[str, int] = {}
    uso_secundaria = False
    grupo: str | None = None

    def _de_secundaria(i: int) -> None:
        nonlocal uso_secundaria
        if grupo == _GRUPO_VIGENTE and i < len(secundaria):
            sub = _texto(secundaria[i])
            if sub:
                columnas[sub] = i
                uso_secundaria = True

    for i, valor in enumerate(primaria):
        texto = _texto(valor)
        if not texto:
            _de_secundaria(i)
            continue
        mayusculas = texto.upper()
        if mayusculas in _GRUPOS_CABECERA:
            grupo = mayusculas
            _de_secundaria(i)
            continue
        grupo = None
        columnas[texto] = i
    return columnas, uso_secundaria


def _cabecera(filas: list[tuple[Any, ...]]) -> tuple[dict[str, int], int]:
    """Ubica la fila de cabecera y devuelve `(columnas, primera_fila_de_datos)`.

    No se asume en una posicion fija: la hoja suelta la trae en la primera
    fila, el reporte completo un par de filas mas abajo, detras de su titulo.
    """
    limite = min(len(filas), _TOPE_FILAS_CABECERA)
    for indice in range(limite):
        siguiente = filas[indice + 1] if indice + 1 < len(filas) else ()
        columnas, uso_secundaria = _combinar_cabecera(filas[indice], siguiente)
        if all(o in columnas for o in _OBLIGATORIAS):
            return columnas, indice + (2 if uso_secundaria else 1)
    raise ErrorValidacion(
        f"Al archivo le falta alguna de estas columnas: {', '.join(_OBLIGATORIAS)}. "
        f"Se esperaba la estructura del distributivo que exporta el sistema academico.",
        campo="archivo",
    )


class CarreraAgregada(Exception):
    """La fila junta varias carreras o sedes en una celda."""


#: Firma de un documento OLE2, que es lo que realmente son los `.xls` que
#: exporta el sistema academico. Se mira el contenido y no la extension porque
#: el origen tambien entrega `.xlsx` con nombre `.xls`.
_FIRMA_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


#: Nombre de la hoja que trae la carga horaria en el reporte completo del PAO.
_HOJA_DISTRIBUTIVO = "distributivo"


def _elegir_hoja(nombres: Sequence[str], hoja: str | None) -> str:
    """Sin hoja explicita, busca `Distributivo` por nombre antes de asumir la primera.

    El reporte completo trae diez pestanas y `Distributivo` no es la primera
    (`Antecedentes` lo es); adivinar por posicion la perderia.
    """
    if hoja:
        return hoja
    for nombre in nombres:
        if nombre.strip().lower() == _HOJA_DISTRIBUTIVO:
            return nombre
    return nombres[0]


def _filas_xlsx(datos: bytes, hoja: str | None) -> Iterator[tuple[Any, ...]]:
    libro = load_workbook(BytesIO(datos), read_only=True, data_only=True)
    try:
        pagina = libro[_elegir_hoja(libro.sheetnames, hoja)]
        yield from pagina.iter_rows(values_only=True)
    finally:
        libro.close()


def _filas_xls(datos: bytes, hoja: str | None) -> Iterator[tuple[Any, ...]]:
    """Formato Excel 97-2003, que openpyxl no lee.

    El sistema academico exporta en este formato, asi que convertirlo a mano
    antes de cargarlo seria un paso manual en cada importacion.
    """
    libro = xlrd.open_workbook(file_contents=datos)
    pagina = libro.sheet_by_name(_elegir_hoja(libro.sheet_names(), hoja))
    for numero in range(pagina.nrows):
        yield tuple(pagina.row_values(numero))


def _filas_del_libro(origen: str | Path | bytes, hoja: str | None) -> Iterator[tuple[Any, ...]]:
    if isinstance(origen, bytes):
        datos = origen
    else:
        camino = Path(origen)
        if not camino.exists():
            raise ErrorValidacion(f"El archivo no existe: {camino}", campo="archivo")
        datos = camino.read_bytes()

    if not datos:
        raise ErrorValidacion("El archivo esta vacio", campo="archivo")

    try:
        if datos[:8] == _FIRMA_OLE2:
            yield from _filas_xls(datos, hoja)
        else:
            yield from _filas_xlsx(datos, hoja)
    except ErrorValidacion:
        raise
    except Exception as exc:  # el motivo lo pone la libreria que abre el libro
        raise ErrorValidacion(
            f"No se pudo leer el archivo como hoja de calculo: {exc}", campo="archivo"
        ) from exc


class LectorPaoExcel:
    """Convierte el distributivo del sistema academico en filas planas."""

    def leer(
        self,
        origen: str | Path | bytes,
        *,
        pao: str,
        interciclo: bool = False,
        hoja: str | None = None,
    ) -> tuple[list[FilaCrudaDistributivo], list[tuple[int, str, str]]]:
        """Devuelve `(filas, rechazadas)`.

        `origen` es una ruta o el contenido del archivo. Acepta bytes para que
        la carga desde la interfaz no tenga que escribir un temporal en disco.

        `pao` es el codigo del periodo —`262651`, `261650`—, que el archivo no
        trae. Cada rechazada es `(numero de fila, identificacion, motivo)`.
        """
        # Se materializa: la cabecera puede estar a un par de filas de
        # distancia y hay que poder mirar hacia adelante para ubicarla. Un PAO
        # entero cabe de sobra en memoria — el mas grande cargado pesa 820 KB.
        libro = list(_filas_del_libro(origen, hoja))
        if not libro:
            raise ErrorValidacion("El archivo esta vacio", campo="archivo")

        col, inicio = _cabecera(libro)
        columna_identificacion = col["Identificación"]

        filas: list[FilaCrudaDistributivo] = []
        rechazadas: list[tuple[int, str, str]] = []

        for numero, cruda in enumerate(libro[inicio:], start=inicio + 1):
            identificacion = (
                _texto(cruda[columna_identificacion]) if columna_identificacion < len(cruda) else ""
            )
            if not identificacion:
                continue
            try:
                filas.append(self._fila(numero, cruda, col, pao, interciclo))
            except CarreraAgregada as exc:
                rechazadas.append((numero, identificacion, str(exc)))

        log.info(
            "Distributivo del sistema academico leido",
            extra={"filas": len(filas), "rechazadas": len(rechazadas), "pao": pao},
        )
        return filas, rechazadas

    # ------------------------------------------------------------------ fila
    def _fila(
        self,
        numero: int,
        cruda: tuple[Any, ...],
        col: dict[str, int],
        pao: str,
        interciclo: bool,
    ) -> FilaCrudaDistributivo:
        def valor(nombre: str) -> str:
            posicion = col.get(nombre)
            if posicion is None or posicion >= len(cruda):
                return ""
            return _texto(cruda[posicion])

        carrera, sede = valor("Carrera/Programa"), valor("Sede")
        for campo, contenido in (("carrera", carrera), ("sede", sede)):
            if "," in contenido:
                raise CarreraAgregada(
                    f"La {campo} junta varios valores en una celda: '{contenido[:60]}'"
                )

        facultad = valor("Facultad").upper()
        horas = {clave: self._numero(cruda, col, clave) for clave in _CLAVES_HORAS if clave in col}

        return FilaCrudaDistributivo(
            numero_fila=numero,
            identificacion=valor("Identificación"),
            pao=pao,
            facultad=FACULTADES.get(facultad, facultad),
            carrera=self._carrera(carrera, sede, valor("Nivel"), valor("Modalidad")),
            sede=_o_nada(sede),
            nombre=valor("Apellidos y Nombres"),
            titularidad=_o_nada(valor("Titularidad")),
            dedicacion=_o_nada(valor("Dedicación")),
            categoria=self._categoria(valor("Categoría")),
            nivel=_o_nada(valor("Nivel")),
            medida=_o_nada(valor("Medida")),
            horas=horas,
            interciclo=interciclo,
            sistema=DatosSistemaAcademico(
                estado_validacion=ESTADOS.get(valor("Estado de la Validación").upper()),
                fase=_o_nada(valor("Fase")),
                semanas=self._entero(valor("N.º Semanas")),
                relacion_laboral=_o_nada(valor("Relación Laboral")),
                tutor_posgrado=valor("Tutores Posgrado").upper() in {"SÍ", "SI"},
                tutor_medicina=valor("Tutores Medicina").upper() in {"SÍ", "SI"},
            ),
        )

    @staticmethod
    def _categoria(valor: str) -> str | None:
        limpio = _o_nada(valor)
        return CATEGORIAS.get((limpio or "").upper(), limpio) if limpio else None

    @staticmethod
    def _carrera(carrera: str, sede: str, nivel: str, modalidad: str) -> str | None:
        """Recompone `SEDE:NOMBRE - NIVEL - MODALIDAD`, como el consolidado.

        Sin prefijo de sede conocido se deja el nombre solo: es lo que hace el
        consolidado con las carreras que no lo llevan, como `MEDICINA (R)`.
        """
        if not carrera or carrera.upper() in _AUSENTES:
            return None
        prefijo = PREFIJO_SEDE.get(sede.upper())
        partes = [p for p in (nivel, modalidad) if p and p.upper() not in _AUSENTES]
        nombre = f"{prefijo}:{carrera}" if prefijo else carrera
        return " - ".join([nombre, *partes]) if partes else nombre

    @staticmethod
    def _numero(cruda: tuple[Any, ...], col: dict[str, int], clave: str) -> float:
        try:
            return float(cruda[col[clave]] or 0)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _entero(valor: str) -> int | None:
        try:
            return int(float(valor))
        except (TypeError, ValueError):
            return None
